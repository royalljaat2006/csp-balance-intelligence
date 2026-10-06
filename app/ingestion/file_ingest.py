"""
Shared file-ingestion orchestration for transaction .xlsx files — validate,
parse+upsert (transaction_ingest.ingest_workbook), file to processed/failed,
and record an IngestLog. The ONE place this happens: manage.py
ingest_transactions (folder-watch) and manage.py poll_telegram both call
this — never a second, source-specific copy of the validate/ingest/file
sequence. `source`/`external_ref`/`source_hash` only label the IngestLog
row for the audit trail; they change no ingestion behaviour.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import structlog
from common.cache import bump_generation
from csp.models import IngestLog
from django.utils import timezone

from .file_safety import IngestionPaths, move_to_failed, move_to_processed
from .transaction_ingest import ingest_workbook
from .xlsx_validation import (
    validate_file_basics,
    validate_required_columns,
    validate_required_sheet,
)

logger = structlog.get_logger("ingestion")


def ingest_transaction_file(
    paths: IngestionPaths,
    file_path: Path,
    *,
    source: str,
    external_ref: str = "",
    source_hash: str = "",
) -> IngestLog:
    """Validates, parses, and upserts one transaction .xlsx file already
    sitting in `paths.incoming`, then files it to processed/ or failed/ —
    never deleted, never left in incoming/. Returns the finished IngestLog
    row (never raises for an ingestion-content problem — that's what
    log.status/error_summary communicate)."""
    log = IngestLog.objects.create(
        source=source,
        file_name=file_path.name,
        external_ref=external_ref,
        source_hash=source_hash,
        status=IngestLog.Status.FAILED,  # default to failed; flipped below on success paths
    )
    job = logger.bind(job_id=log.id, file=file_path.name, source=source)
    job.info("started", stage="INGESTION")

    basics = validate_file_basics(file_path)
    if not basics.ok:
        move_to_failed(paths, file_path)
        job.warning("rejected", stage="VALIDATION", reason=basics.reason)
        _fail(log, basics.reason)
        return log

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — any corrupt/unreadable file is a rejection, not a crash
        move_to_failed(paths, file_path)
        reason = f"Rejected {file_path.name}: could not open workbook ({exc})"
        job.warning("rejected", stage="VALIDATION", reason="unreadable_workbook")
        _fail(log, reason)
        return log

    try:
        validate_required_sheet(wb.sheetnames)
        sheet = wb["Success"]
        header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
        validate_required_columns(header)
    except Exception as exc:  # noqa: BLE001
        wb.close()  # release the file handle *before* moving it — Windows can't rename an open file
        move_to_failed(paths, file_path)
        job.warning("rejected", stage="VALIDATION", reason=str(exc))
        _fail(log, str(exc))
        return log
    wb.close()

    result = ingest_workbook(file_path, header)
    job.info(
        "completed",
        stage="VALIDATION",
        rows_read=result.rows_read,
        rows_valid=result.rows_valid,
        rows_rejected=result.rows_rejected,
    )
    if result.new_csp_codes:
        job.info(
            "stub_csps_created",
            stage="DATABASE",
            count=len(result.new_csp_codes),
            codes=sorted(result.new_csp_codes)[:20],  # cap — this is a log line, not a dump
        )
    job.info("upserted", stage="DATABASE", rows_upserted=result.rows_upserted)
    bump_generation()  # new transaction rows landed -> invalidate cached dashboard rollups

    move_to_processed(paths, file_path)
    log.status = (
        IngestLog.Status.SUCCESS if result.rows_rejected == 0 else IngestLog.Status.PARTIAL
    )
    log.rows_read = result.rows_read
    log.rows_valid = result.rows_valid
    log.rows_rejected = result.rows_rejected
    log.rows_upserted = result.rows_upserted
    log.finished_at = timezone.now()
    if result.rows_rejected:
        log.error_summary = (
            f"{result.rows_rejected} in-scope row(s) had a malformed required field "
            "(amount, date, reference number, or CSP code) and were skipped."
        )
    if result.new_csp_codes:
        note = (
            f"{len(result.new_csp_codes)} CSP code(s) not yet in the CALLING SHEET "
            "were stub-created."
        )
        log.error_summary = f"{log.error_summary} {note}".strip()
    log.save(
        update_fields=[
            "status",
            "rows_read",
            "rows_valid",
            "rows_rejected",
            "rows_upserted",
            "finished_at",
            "error_summary",
        ]
    )
    job.info("finished", stage="INGESTION", status=log.status)
    return log


def _fail(log: IngestLog, message: str) -> None:
    log.error_summary = message
    log.finished_at = timezone.now()
    log.save(update_fields=["status", "error_summary", "finished_at"])
