"""
Statement file parsing with content-based format detection.

Observed (2026-10-12 sample, SBI pool account): a file named `*.xls` that is
really UTF-8 tab-delimited text -- a header block of `Label :<TAB>value` lines,
then a `Txn Date ...` table, then a `**This is a computer generated ...`
footer. Never trust the extension: `detect_format()` looks at the bytes.

Supported: tab-delimited text (observed) and genuine `.xlsx` (same layout,
read through openpyxl). Recognised but NOT supported (reported as a clear
StatementFormatError so the file is recorded and quarantined, never
mis-parsed): legacy binary `.xls` (OLE2) and HTML-as-xls.

All money is parsed to Decimal; Indian digit grouping (`57,68,774.00`) is
handled. A row that cannot be parsed is returned in `malformed`, never
silently dropped or coerced.
"""

from __future__ import annotations

import datetime as dt
import decimal
import io
import re
from dataclasses import dataclass, field

Decimal = decimal.Decimal

OLE2_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
ZIP_MAGIC = b"PK\x03\x04"

REQUIRED_HEADERS = [
    "Txn Date",
    "Value Date",
    "Description",
    "Ref No./Cheque No.",
    "Branch Code",
    "Debit",
    "Credit",
    "Balance",
]


class StatementFormatError(ValueError):
    """The file is not a statement layout this parser can read."""


@dataclass(frozen=True)
class ParsedRow:
    index: int  # 1-based position in the table, in file order
    txn_date: dt.date
    value_date: dt.date | None
    description: str
    ref_column: str
    branch_code: str
    debit: Decimal
    credit: Decimal
    balance: Decimal


@dataclass
class ParsedStatement:
    detected_format: str
    metadata: dict[str, str] = field(default_factory=dict)
    account_number: str = ""
    start_date: dt.date | None = None
    end_date: dt.date | None = None
    opening_balance: Decimal | None = None
    opening_date: dt.date | None = None
    book_balance: Decimal | None = None
    available_balance: Decimal | None = None
    hold_value: Decimal | None = None
    rows: list[ParsedRow] = field(default_factory=list)
    malformed: list[tuple[int, str]] = field(default_factory=list)
    footer_seen: bool = False


def detect_format(data: bytes) -> str:
    head = data[:4096]
    if head.startswith(ZIP_MAGIC):
        return "xlsx"
    if head.startswith(OLE2_MAGIC):
        return "xls_binary"
    if b"\x00" in head:
        return "binary_unknown"
    text = head.decode("utf-8-sig", errors="replace").lstrip().lower()
    if text.startswith("<"):
        return "html"
    # Statement text: header block followed by a tab-delimited Txn Date table.
    full = _decode(data)
    for line in full.splitlines():
        if line.startswith("Txn Date") and "\t" in line:
            return "tsv_text"
    return "unrecognised"


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


CENT = Decimal("0.01")


def parse_money(raw: str | None) -> Decimal | None:
    """'57,68,774.00' -> Decimal('5768774.00'); '' -> None; 'x' -> ValueError.
    A trailing Cr/Dr is honoured (Dr -> negative). Surrounding quotes (added when a
    statement is re-saved from a spreadsheet) are stripped. Results are always
    quantized to 2 places so '1000' and '1000.00' are the same value everywhere
    (row identity, reconciliation, storage)."""
    if raw is None:
        return None
    s = raw.strip().strip('"').strip("'").replace(",", "").replace("₹", "").replace(" ", "")
    if not s:
        return None
    sign = 1
    low = s.lower()
    if low.endswith("dr"):
        sign, s = -1, s[:-2]
    elif low.endswith("cr"):
        s = s[:-2]
    try:
        return (sign * Decimal(s)).quantize(CENT, rounding=decimal.ROUND_HALF_EVEN)
    except decimal.InvalidOperation as exc:
        raise ValueError(f"not a money value: {raw!r}") from exc


_DATE_FORMATS = (
    "%d/%m/%Y",
    "%d %b %Y",
    "%d-%b-%Y",
    "%d-%b-%y",
    "%d-%m-%Y",
    "%Y-%m-%d",
)


def parse_date(raw: str | None) -> dt.date | None:
    if raw is None:
        return None
    s = raw.strip().strip('"').strip("'").strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"not a date: {raw!r}")


def _grid_from_text(text: str) -> list[list[str]]:
    return [line.split("\t") for line in text.splitlines()]


def _grid_from_xlsx(data: bytes) -> list[list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        grid: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells: list[str] = []
            for v in row:
                if v is None:
                    cells.append("")
                elif isinstance(v, dt.datetime):
                    cells.append(v.strftime("%d/%m/%Y"))
                elif isinstance(v, dt.date):
                    cells.append(v.strftime("%d/%m/%Y"))
                else:
                    cells.append(str(v))
            grid.append(cells)
        return grid
    finally:
        wb.close()


def peek_header(head: bytes) -> ParsedStatement:
    """Cheap, read-only look at the first bytes of a file: real format plus (for
    tab-delimited statements) the header fields -- account, period, opening and
    book balance. Never parses the transaction table."""
    fmt = detect_format(head)
    result = ParsedStatement(detected_format=fmt)
    if fmt != "tsv_text":
        return result
    grid = _grid_from_text(_decode(head))
    idx = next(
        (i for i, cells in enumerate(grid) if cells and cells[0].strip() == "Txn Date"), None
    )
    if idx is not None:
        _parse_metadata(grid[:idx], result)
    return result


def _unquote(cell: str) -> str:
    """Undo CSV-style quoting added when a statement is re-saved from a spreadsheet: a cell
    that contains a quote character comes back wrapped in quotes (inner quotes doubled).
    Only a cell that BOTH starts and ends with a quote is touched."""
    if len(cell) >= 2 and cell.startswith('"') and cell.endswith('"'):
        return cell[1:-1].replace('""', '"')
    return cell


def _collapse(label: str) -> str:
    return re.sub(r"\s+", " ", label).strip()


def parse_statement(data: bytes) -> ParsedStatement:
    fmt = detect_format(data)
    if fmt == "tsv_text":
        grid = _grid_from_text(_decode(data))
    elif fmt == "xlsx":
        grid = _grid_from_xlsx(data)
    elif fmt == "xls_binary":
        raise StatementFormatError(
            "Legacy binary .xls (OLE2) is not supported -- re-export as text/.xlsx "
            "(no xlrd dependency is installed)."
        )
    elif fmt == "html":
        raise StatementFormatError("HTML-formatted statement is not supported.")
    else:
        raise StatementFormatError(f"Unrecognised statement format ({fmt}).")
    return _parse_grid(grid, fmt)


def _parse_grid(grid: list[list[str]], fmt: str) -> ParsedStatement:
    header_idx = next(
        (i for i, cells in enumerate(grid) if cells and cells[0].strip() == "Txn Date"), None
    )
    if header_idx is None:
        raise StatementFormatError("No 'Txn Date' table header found.")

    result = ParsedStatement(detected_format=fmt)
    _parse_metadata(grid[:header_idx], result)

    header = [c.strip() for c in grid[header_idx]]
    missing = [h for h in REQUIRED_HEADERS if h not in header]
    if missing:
        raise StatementFormatError(f"Transaction table is missing columns: {missing}")
    col = {name: header.index(name) for name in REQUIRED_HEADERS}
    width = max(col.values()) + 1

    index = 0
    for lineno, cells in enumerate(grid[header_idx + 1 :], start=header_idx + 2):
        first = cells[0].strip() if cells else ""
        if first.startswith("**"):
            result.footer_seen = True
            break
        if not any(c.strip() for c in cells):
            continue
        index += 1
        if len(cells) < width:
            result.malformed.append((index, f"line {lineno}: {len(cells)} columns, need {width}"))
            continue
        try:
            txn_date = parse_date(cells[col["Txn Date"]])
            if txn_date is None:
                raise ValueError("missing Txn Date")
            debit = parse_money(cells[col["Debit"]]) or Decimal("0.00")
            credit = parse_money(cells[col["Credit"]]) or Decimal("0.00")
            balance = parse_money(cells[col["Balance"]])
            if balance is None:
                raise ValueError("missing Balance")
            if debit < 0 or credit < 0:
                raise ValueError("negative debit/credit")
            if debit and credit:
                raise ValueError("both debit and credit present")
            result.rows.append(
                ParsedRow(
                    index=index,
                    txn_date=txn_date,
                    value_date=parse_date(cells[col["Value Date"]]),
                    description=_unquote(cells[col["Description"]]),
                    ref_column=_unquote(cells[col["Ref No./Cheque No."]]),
                    branch_code=cells[col["Branch Code"]].strip(),
                    debit=debit,
                    credit=credit,
                    balance=balance,
                )
            )
        except ValueError as exc:
            result.malformed.append((index, f"line {lineno}: {exc}"))
    return result


def _parse_metadata(lines: list[list[str]], result: ParsedStatement) -> None:
    for cells in lines:
        if not cells:
            continue
        first = cells[0]
        if ":" in first and first.split(":", 1)[1].strip():
            label, value = first.split(":", 1)  # single-cell "Label : value"
        elif ":" in first:
            label = first.split(":", 1)[0]
            value = next((c for c in cells[1:] if c.strip()), "")
        else:
            continue
        label, value = _collapse(label), value.strip()
        if not label:
            continue
        result.metadata[label] = value
        try:
            _apply_metadata(label, value, result)
        except ValueError:
            # A header field we cannot parse is left unset (and surfaces as a
            # reconciliation note) -- it must not abort the whole statement.
            continue


def _apply_metadata(label: str, value: str, result: ParsedStatement) -> None:
    if label == "Account Number":
        result.account_number = value.lstrip("_").strip()
    elif label == "Start Date":
        result.start_date = parse_date(value)
    elif label == "End Date":
        result.end_date = parse_date(value)
    elif label.startswith("Opening Balance"):
        result.opening_balance = parse_money(value)
        m = re.search(r"as on\s+(.+)$", label, flags=re.IGNORECASE)
        if m:
            result.opening_date = parse_date(m.group(1))
    elif label == "Book Balance":
        result.book_balance = parse_money(value)
    elif label == "Available Balance":
        result.available_balance = parse_money(value)
    elif label == "Hold Value":
        result.hold_value = parse_money(value)
