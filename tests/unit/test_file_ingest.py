"""
ingest_transaction_file — the shared validate/parse/upsert/file orchestrator
behind both manage.py ingest_transactions (folder-watch) and manage.py
poll_telegram (Telegram-delivered). ingest_workbook's own row-level
behaviour is covered by test_transaction_ingest.py; this file is about the
orchestration around it: file safety (moved, never deleted), IngestLog
status/field population, and source/external_ref/source_hash labeling.
"""

import openpyxl
import pytest
from csp.models import IngestLog, Transaction
from ingestion.file_ingest import ingest_transaction_file
from ingestion.file_safety import IngestionPaths

HEADER = [
    "KO ID", "Transaction Date & Time", "Reference Number", "Type of Transaction",
    "From Account", "To Account", "Amount", "Status", "New Date", "CSP Code New",
]


def _make_workbook(paths, rows, name="Transaction Test.xlsx"):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Success"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    path = paths.incoming / name
    wb.save(path)
    return path


def _row(ref="REF001", txn_type="AEPS ONUS Withdrawal", amount="2000", csp="1A850001"):
    return [
        csp, "01-08-2026 08:00:00 AM", ref, txn_type,
        "XXXXXX11111", "XXXXXX71556", amount, "Success", "01-08-2026", csp,
    ]


@pytest.fixture
def paths(tmp_path):
    p = IngestionPaths(root=tmp_path)
    p.ensure_exist()
    return p


@pytest.mark.django_db
def test_success_moves_file_and_records_ingest_log(paths):
    file_path = _make_workbook(paths, [_row()])
    log = ingest_transaction_file(paths, file_path, source="transactions")

    assert log.status == IngestLog.Status.SUCCESS
    assert log.source == "transactions"
    assert log.rows_upserted == 1
    assert not file_path.exists()
    assert len(list(paths.processed.glob("*.xlsx"))) == 1
    assert Transaction.objects.filter(ref_number="REF001").exists()


@pytest.mark.django_db
def test_labels_source_external_ref_and_source_hash(paths):
    file_path = _make_workbook(paths, [_row()])
    log = ingest_transaction_file(
        paths, file_path, source="telegram", external_ref="123", source_hash="UNIQUE_ABC"
    )
    assert log.source == "telegram"
    assert log.external_ref == "123"
    assert log.source_hash == "UNIQUE_ABC"


@pytest.mark.django_db
def test_rows_rejected_yields_partial_status(paths):
    file_path = _make_workbook(paths, [_row(ref="OK1"), _row(ref="BAD1", amount="not-a-number")])
    log = ingest_transaction_file(paths, file_path, source="transactions")

    assert log.status == IngestLog.Status.PARTIAL
    assert log.rows_rejected == 1
    assert "malformed required field" in log.error_summary


@pytest.mark.django_db
def test_missing_required_sheet_is_failed_and_moved_to_failed_dir(paths):
    wb = openpyxl.Workbook()
    wb.active.title = "NotSuccess"
    file_path = paths.incoming / "bad_sheet.xlsx"
    wb.save(file_path)

    log = ingest_transaction_file(paths, file_path, source="transactions")

    assert log.status == IngestLog.Status.FAILED
    assert not file_path.exists()
    assert len(list(paths.failed.glob("*.xlsx"))) == 1
    assert "Success" in log.error_summary


@pytest.mark.django_db
def test_missing_required_columns_is_failed(paths):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Success"
    sheet.append(["Wrong", "Columns"])
    file_path = paths.incoming / "bad_columns.xlsx"
    wb.save(file_path)

    log = ingest_transaction_file(paths, file_path, source="transactions")

    assert log.status == IngestLog.Status.FAILED
    assert len(list(paths.failed.glob("*.xlsx"))) == 1


@pytest.mark.django_db
def test_corrupt_workbook_is_failed_not_crashed(paths):
    file_path = paths.incoming / "corrupt.xlsx"
    file_path.write_bytes(b"this is not a real xlsx file")

    log = ingest_transaction_file(paths, file_path, source="transactions")

    assert log.status == IngestLog.Status.FAILED
    assert "could not open workbook" in log.error_summary
    assert not file_path.exists()
    assert len(list(paths.failed.glob("*.xlsx"))) == 1


@pytest.mark.django_db
def test_empty_file_is_rejected_by_basic_validation(paths):
    file_path = paths.incoming / "empty.xlsx"
    file_path.write_bytes(b"")

    log = ingest_transaction_file(paths, file_path, source="transactions")

    assert log.status == IngestLog.Status.FAILED
    assert "empty" in log.error_summary
    assert len(list(paths.failed.glob("*.xlsx"))) == 1


@pytest.mark.django_db
def test_new_csp_codes_are_noted_in_error_summary(paths):
    file_path = _make_workbook(paths, [_row(csp="9Z999999")])
    log = ingest_transaction_file(paths, file_path, source="transactions")

    assert log.status == IngestLog.Status.SUCCESS
    assert "stub-created" in log.error_summary
