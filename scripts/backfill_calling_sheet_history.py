"""
One-off backfill: historical DailyBalance rows from archived Calling Sheet
workbooks (e.g. a team's own periodic backup folder on the R730 —
confirmed source: /home/shumit/sbikisok/_calling_sheet_backups/, files
named "Calling Sheet_YYYYMMDD_HHMMSS.xlsx").

Why: the live `ingest_calling_sheet` command always writes to *today's*
date only, so a freshly-deployed instance has no balance history before
its own go-live -- "today vs yesterday" and trailing-N-day trend widgets
read as empty/zero until enough days accumulate naturally. This backfills
that gap for real from an existing archive, instead of waiting.

Reuses the exact same parsing logic as the live command
(ingestion/calling_sheet_ingest.parse_sheet_values). Confirmed by
inspection against the real archive (2026-10-07): only two column headers
differ from the live sheet's export ("CSP Code" vs "CSP ID", "Mobile
Number" vs "CSP Mobile number"), renamed below before parsing -- currency/
count cell formatting is otherwise identical, so no other adaptation is
needed. A handful of the very earliest files in that archive predate a
sheet restructuring and are missing required columns entirely; these are
skipped with a message, not treated as a crash.

Deliberately does NOT touch Csp master-data fields for existing CSPs
(name/email/mobile/etc.) -- only creates a bare stub, tagged
status="historical_only", if a historical file references a csp_code not
already known, so today's live-ingested master data is never regressed by
an older snapshot. ingest_calling_sheet clears that tag automatically the
moment a live poll actually sees the CSP again (see its _CSP_UPDATE_FIELDS
and Csp(status="") in that file). Only DailyBalance is fully upserted,
keyed on (csp, balance_date) same as the live command, tagged
source=HISTORY (not CALLING_SHEET) to distinguish backfilled rows from
live polls in the data itself.

Skips today's own date -- today's DailyBalance already comes from the live
60s-interval poll, which is fresher than a one-off backup snapshot.

Usage (run from inside a worker container that has the TRANSACTION_DATA_DIR
bind mount and DB access -- e.g. worker-calling-sheet):
    1. Copy the archived .xlsx files into a subfolder *inside* that bind
       mount, but NOT into incoming/ (ingest_transactions' watcher globs
       incoming/*.xlsx indiscriminately and will wrongly try to ingest them
       as transaction files). A sibling folder such as .../data/cs_backfill/
       is safe.
    2. Copy this script alongside them.
    3. docker compose exec -T <worker-service> python3 \
           /data/cs_backfill/backfill_calling_sheet_history.py
    4. Re-run dbt + sync_monthly_summary afterward so marts reflect the
       newly-backfilled history.
"""
import datetime as dt
import os
import re
import sys
from pathlib import Path

import django

sys.path.insert(0, "/app/app")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings.production")
django.setup()

import openpyxl  # noqa: E402
from csp.models import Csp, DailyBalance  # noqa: E402
from django.utils import timezone  # noqa: E402
from ingestion.calling_sheet_ingest import (  # noqa: E402
    CallingSheetValidationError,
    parse_sheet_values,
)

BACKFILL_DIR = Path(os.environ.get("CS_BACKFILL_DIR", "/data/cs_backfill"))
HEADER_RENAMES = {
    "CSP Code": "CSP ID",
    "Mobile Number": "CSP Mobile number",
}
FILENAME_RE = re.compile(r"Calling Sheet_(\d{8})_(\d{6})\.xlsx")

today = timezone.localdate()


def parse_file_date(path: Path) -> dt.date | None:
    m = FILENAME_RE.match(path.name)
    if not m:
        return None
    return dt.datetime.strptime(m.group(1), "%Y%m%d").date()


def load_all_values(path: Path) -> list[list[str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Calling Sheet New"] if "Calling Sheet New" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])
    wb.close()
    return rows


def main():
    dated = [
        (file_date, p)
        for p in BACKFILL_DIR.glob("Calling Sheet_*.xlsx")
        if (file_date := parse_file_date(p)) is not None and file_date < today
    ]
    dated.sort(key=lambda pair: pair[0])
    print(f"Found {len(dated)} dated backup files before {today}")

    existing_codes = set(Csp.objects.values_list("csp_code", flat=True))
    total_days = 0
    total_balance_rows = 0

    for file_date, path in dated:
        try:
            all_values = load_all_values(path)
        except Exception as exc:  # noqa: BLE001
            print(f"SKIP {path.name}: could not read workbook: {exc}")
            continue

        if len(all_values) >= 2:
            header = all_values[1]
            all_values[1] = [HEADER_RENAMES.get(h, h) for h in header]

        try:
            result = parse_sheet_values(all_values)
        except CallingSheetValidationError as exc:
            print(f"SKIP {path.name}: {exc}")
            continue

        # status="historical_only" keeps these out of "CSPs tracked" (see
        # csp/services.py's get_overview) -- they exist only to satisfy
        # DailyBalance's FK for a CSP seen in an old snapshot, not because
        # they're part of the current roster.
        new_codes = {r.csp_code for r in result.rows} - existing_codes
        if new_codes:
            Csp.objects.bulk_create(
                [
                    Csp(csp_code=c, last_seen_date=file_date, status="historical_only")
                    for c in new_codes
                ],
                ignore_conflicts=True,
            )
            existing_codes |= new_codes

        # Dedupe by csp_code within this one file/date: Postgres's ON CONFLICT
        # DO UPDATE cannot affect the same conflict target twice in one
        # statement, and a handful of these sheets have a CSP listed more
        # than once on the same day. Last occurrence in the sheet wins.
        rows_by_csp = {
            row.csp_code: row for row in result.rows if row.avg_balance is not None
        }
        balance_objs = [
            DailyBalance(
                csp_id=row.csp_code,
                balance_date=file_date,
                daily_avg_balance=row.avg_balance,
                account_count=row.account_count,
                source=DailyBalance.Source.HISTORY,
            )
            for row in rows_by_csp.values()
        ]
        if balance_objs:
            DailyBalance.objects.bulk_create(
                balance_objs,
                update_conflicts=True,
                unique_fields=["csp", "balance_date"],
                update_fields=["daily_avg_balance", "account_count", "source"],
                batch_size=1000,
            )
        total_days += 1
        total_balance_rows += len(balance_objs)
        print(f"{path.name} -> {file_date}: {len(balance_objs)} balance rows upserted")

    distinct_dates = DailyBalance.objects.values_list("balance_date", flat=True).distinct().count()
    print(
        f"\nDone. Processed {total_days} files, {total_balance_rows} total balance-row "
        f"upserts. DailyBalance now has data for {distinct_dates} distinct dates."
    )


if __name__ == "__main__":
    main()
